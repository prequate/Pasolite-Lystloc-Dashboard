#!/usr/bin/env python3
"""
Pasolite - LystLoc Weekly Check-ins Report cleaner (fresh-start version).

Input : GWC104 "Weekly Checkins Report" raw export from LystLoc.
Output: same S.No -> Checkout Location columns (S.No/Date forward-filled so
        every row stands alone), plus Checkout Meeting Notes split into its
        7 constituent fields as separate columns.

This is meant to become the new standard LystLoc intake step (replacing the
old, unstructured handling) feeding the Weekly Workbook / Revenue Scorecard
process (Task A in Pasolite Instructions.docx). This script only cleans the
raw check-ins export; it does not touch the master workbook itself.
"""
import re
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


NOTES_PATTERN = re.compile(
    r"Lead\s*:\s*(?P<lead>.*?)\s*,\s*"
    r"Estimated Project value\s*:\s*(?P<value>.*?)\s*,\s*"
    r"Project Requirement Timeline\s*:\s*(?P<timeline>.*?)\s*,\s*"
    r"Client Satisfaction\s*\(1-3\)\s*:\s*(?P<satisfaction>.*?)\s*,\s*"
    r"Next Follow-up Date\s*:\s*(?P<followup>.*?)\s*,\s*"
    r"Type of Lead\s*:\s*(?P<leadtype>.*?)\s*,\s*"
    r"Remarks\s*:\s*(?P<remarks>.*)$",
    re.DOTALL,
)


def clean(v):
    if v is None:
        return ""
    return str(v).strip()


# Estimated Project Value, Project Requirement Timeline and Client
# Satisfaction all come from LystLoc as "<code> - <label>" (e.g.
# "1 - Below  ₹50,000", "4 - 1 to 2 Months", "2 -Interested"). Harsh only
# wants the label, not the numeric code, so strip a leading "<digit(s)>
# <spaces> - <spaces>" and collapse any internal double-spacing.
CODE_PREFIX = re.compile(r"^\s*\d+\s*-\s*")


def strip_code_prefix(v):
    v = CODE_PREFIX.sub("", v).strip()
    v = re.sub(r"\s{2,}", " ", v)
    return v


def split_notes(notes):
    """Return the 7 parsed fields for a Checkout Meeting Notes cell, or
    7 blanks if there's nothing to parse (a plain '-' / empty cell)."""
    text = clean(notes)
    if text in ("", "-"):
        return ["", "", "", "", "", "", ""]
    m = NOTES_PATTERN.search(text)
    if not m:
        # Should not happen (all 104 real rows matched during dry-run), but
        # never silently drop data: surface the raw text in Remarks instead.
        return ["", "", "", "", "", "", f"[UNPARSED] {text}"]
    d = m.groupdict()
    remarks = d["remarks"].strip().strip(",").strip()
    return [
        d["lead"].strip(),
        strip_code_prefix(d["value"].strip()),
        strip_code_prefix(d["timeline"].strip()),
        strip_code_prefix(d["satisfaction"].strip()),
        d["followup"].strip(),
        d["leadtype"].strip(),
        remarks,
    ]


def clean_one(src_path, out_path, source_label):
    wb_in = openpyxl.load_workbook(src_path, data_only=True)
    ws_in = wb_in["Detailed Report"]

    header = [c.value for c in ws_in[2]]
    assert header == [
        "S.No", "Date", "Name", "Employee Status", "Checkin Time",
        "Checkin Location", "Checkin Meeting Notes", "Checkout Time",
        "Checkout Location", "Checkout Meeting Notes",
    ], f"Unexpected source header, stopping: {header}"

    new_headers = [
        "S.No", "Date", "Name", "Employee Status", "Checkin Time",
        "Checkin Location", "Checkin Meeting Notes", "Checkout Time",
        "Checkout Location",
        "Lead", "Estimated Project Value", "Project Requirement Timeline",
        "Client Satisfaction (1-3)", "Next Follow-up Date", "Type of Lead",
        "Remarks",
    ]

    out_rows = []
    last_sno, last_date = None, None
    unparsed = []
    no_checkout = []
    blank_sno_count = 0

    for i, row in enumerate(ws_in.iter_rows(min_row=3, max_row=ws_in.max_row, values_only=True), start=3):
        if row[0] is None and all(v is None for v in row):
            continue  # fully blank row, skip

        sno, date, name, status, ci_time, ci_loc, ci_notes, co_time, co_loc, co_notes = row

        # Forward-fill S.No / Date on continuation rows (same employee,
        # multiple check-ins that day) so every row is self-contained.
        if sno is not None:
            last_sno = sno
        else:
            blank_sno_count += 1
        if date is not None:
            last_date = date

        if co_time in (None, "-") and co_loc is None and co_notes is None:
            no_checkout.append((last_sno, clean(name)))

        parsed = split_notes(co_notes)
        if parsed[-1].startswith("[UNPARSED]"):
            unparsed.append((last_sno, clean(name), co_notes))

        out_rows.append([
            last_sno, last_date, clean(name), clean(status),
            clean(ci_time), clean(ci_loc), clean(ci_notes),
            clean(co_time), clean(co_loc),
        ] + parsed)

    # ---- write output workbook -------------------------------------------------
    wb_out = openpyxl.Workbook()
    ws = wb_out.active
    ws.title = "LystLoc Checkins - Clean"

    FONT_NAME = "Arial"
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor="F96900")  # Prequate Deep Orange
    body_font = Font(name=FONT_NAME, size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c, h in enumerate(new_headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for r, row in enumerate(out_rows, start=2):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (6, 7, 16)))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(new_headers))}{len(out_rows) + 1}"

    widths = {
        "A": 6, "B": 11, "C": 14, "D": 11, "E": 18, "F": 32, "G": 10,
        "H": 18, "I": 32, "J": 20, "K": 22, "L": 22, "M": 18, "N": 16,
        "O": 24, "P": 34,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 30

    # ---- summary sheet ----------------------------------------------------
    ws2 = wb_out.create_sheet("Read Me")
    ws2.column_dimensions["A"].width = 100
    lines = [
        ("Pasolite - LystLoc Weekly Check-ins, cleaned", True),
        ("", False),
        (f"Source file: GWC104 Weekly Checkins Report, {source_label}.", False),
        (f"Total check-in rows: {len(out_rows)}.", False),
        (f"Rows with a logged lead (Checkout Meeting Notes present): "
         f"{sum(1 for r in out_rows if r[9])}.", False),
        (f"Rows with no lead logged (plain check-in, no notes): "
         f"{sum(1 for r in out_rows if not r[9])}.", False),
        (f"S.No and Date were blank on {blank_sno_count} source rows where LystLoc "
         f"only shows them on the first visit of a group and forward-filled them here so every "
         f"row stands alone for filtering and pivoting.", False),
        ("", False),
        ("Checked in but never checked out (no Checkout Time/Location/Notes):", True),
    ]
    if no_checkout:
        for sno, name in no_checkout:
            lines.append((f"  - S.No {sno}: {name}", False))
    else:
        lines.append(("  - None found.", False))
    lines.append(("", False))
    lines.append(("Rows where Checkout Meeting Notes didn't match the expected "
                   "Lead/Value/Timeline/Satisfaction/Follow-up/Type/Remarks pattern:", True))
    if unparsed:
        for sno, name, raw in unparsed:
            lines.append((f"  - S.No {sno}: {name} -> raw text kept in Remarks column", False))
    else:
        lines.append((f"  - None. All {sum(1 for r in out_rows if r[9])} rows with notes parsed cleanly.", False))

    r = 1
    for text, bold in lines:
        cell = ws2.cell(row=r, column=1, value=text)
        cell.font = Font(name=FONT_NAME, size=11, bold=bold)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    wb_out.move_sheet("Read Me", offset=-1)
    wb_out.active = 0

    wb_out.save(out_path)
    print("Saved:", out_path)
    print("Rows written:", len(out_rows))
    print("No-checkout rows:", no_checkout)
    print("Unparsed rows:", unparsed)
    return out_path


def main():
    # Normally called through run_week.py. Direct use: clean_lystloc.py <raw.xlsx> <out.xlsx> <DD-Mon-YYYY>
    clean_one(sys.argv[1], sys.argv[2], sys.argv[3])


if __name__ == "__main__":
    main()
