#!/usr/bin/env python3
"""
Pasolite weekly LystLoc run (Prequate Advisory).

Usage:  python run_week.py <GWC104 Weekly Checkins Report .xlsx> [--replace]

1. Reads the raw GWC104 export as downloaded from LystLoc. If someone has
   split the notes into columns in Excel (text to columns), the cells are
   joined back first, so either version works.
2. Checks it is one Monday-to-Sunday week and not already loaded
   (use --replace to reload a week on purpose).
3. Cleans it into data/cleaned/ with the standard cleaner.
4. Rebuilds index.html with build_dashboard.py.
5. Prints a short summary, including any drop-down values the dashboard has
   not seen before (for when LystLoc switches to the new options).
Stops with a clear message instead of building anything doubtful.
"""
import json
import os
import subprocess
import sys
import tempfile
from collections import Counter
from datetime import datetime, timedelta

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import clean_lystloc as cl  # noqa: E402

STD = ["S.No", "Date", "Name", "Employee Status", "Checkin Time", "Checkin Location",
       "Checkin Meeting Notes", "Checkout Time", "Checkout Location", "Checkout Meeting Notes"]
CLEANED_DIR = os.path.join(HERE, "data", "cleaned")

KNOWN = {
    "Estimated Project Value": {"Below ₹50,000", "₹50,000 to ₹2 Lakhs", "₹2 Lakhs to ₹5 Lakhs",
                                "₹5 Lakhs to ₹10 Lakhs", "₹10 Lakhs to ₹15 Lakhs", "Above 15 Lakhs"},
    "Project Requirement Timeline": {"Immediate Requirement", "Within 15 Days", "15 Days to 1 Months",
                                     "1 to 2 Months", "Above 2 Months"},
    "Client Satisfaction (1-3)": {"Not Interested", "Interested", "Very Positive / Strong Opportunity ✅"},
    "Type of Lead": {"Architect", "Project", "Builders or Procurement Manager", "Dealer",
                     "Lighting consultant", "Distributor"},
}


def stop(msg):
    print("STOPPED: " + msg)
    sys.exit(2)


def normalise(src):
    """Return a path to a workbook in the standard 10-column GWC104 layout."""
    wb = openpyxl.load_workbook(src, data_only=True)
    if "Detailed Report" not in wb.sheetnames:
        stop(f"no 'Detailed Report' sheet (found {wb.sheetnames}). Is this the GWC104 Weekly Checkins Report?")
    ws = wb["Detailed Report"]
    hdr = [c.value for c in ws[2]]
    if hdr[:10] == STD and all(h is None for h in hdr[10:]) and ws.max_column <= 10:
        return src, "as downloaded"
    if [str(h).strip() if h else h for h in hdr[:9]] != STD[:9]:
        stop(f"unexpected header row: {hdr[:12]}")
    out = openpyxl.Workbook()
    o = out.active
    o.title = "Detailed Report"
    o.append([ws.cell(1, 1).value])
    o.append(STD)
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row is None or all(v is None for v in row):
            o.append([None] * 10)
            continue
        # Excel's text-to-columns turns "₹50,000" into "₹50" and the number 0,
        # so a bare 0 after a split is put back as "000".
        tail = []
        for v in row[9:]:
            if v is None:
                continue
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                v = "000" if v == 0 else (str(int(v)) if float(v).is_integer() else str(v))
            tail.append(str(v))
        o.append(list(row[:9]) + [",".join(tail) if tail else None])
    tmp = os.path.join(tempfile.mkdtemp(), "normalised.xlsx")
    out.save(tmp)
    return tmp, "notes had been split into columns; joined back"


def week_of(src):
    ws = openpyxl.load_workbook(src, data_only=True)["Detailed Report"]
    dates = set()
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row and row[1]:
            dates.add(datetime.strptime(str(row[1]).strip(), "%d-%m-%Y").date())
    if not dates:
        stop("no dates found in the file")
    lo, hi = min(dates), max(dates)
    monday = lo - timedelta(days=lo.weekday())
    sunday = monday + timedelta(days=6)
    if hi > sunday:
        stop(f"file covers more than one week ({lo} to {hi})")
    return monday, sunday, lo, hi


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    replace = "--replace" in sys.argv
    if len(args) != 1:
        stop("give exactly one GWC104 .xlsx file")
    src = args[0]
    norm, layout = normalise(src)
    monday, sunday, lo, hi = week_of(norm)
    label = sunday.strftime("%d-%b-%Y")
    out = os.path.join(CLEANED_DIR, f"Pasolite_LystLoc_Checkins_Cleaned_{label}.xlsx")

    existing = sorted(datetime.strptime(f[-16:-5], "%d-%b-%Y").date()
                      for f in os.listdir(CLEANED_DIR) if f.startswith("Pasolite_LystLoc_Checkins_Cleaned_"))
    if os.path.exists(out) and not replace:
        stop(f"week ending {label} is already loaded. Re-run with --replace to reload it on purpose.")
    gap = None
    if existing and sunday > existing[-1] + timedelta(days=7):
        gap = f"weeks missing between {existing[-1]} and {monday}"

    cl.clean_one(norm, out, label)

    ws = openpyxl.load_workbook(out)["LystLoc Checkins - Clean"]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    forms = [r for r in rows if r.get("Lead")]
    unparsed = [r for r in rows if str(r.get("Remarks") or "").startswith("[UNPARSED]")]
    new_values = {}
    for col, known in KNOWN.items():
        c = Counter(str(r.get(col) or "").strip() for r in forms)
        extra = {k: v for k, v in c.items() if k and k not in known}
        if extra:
            new_values[col] = extra
    fu = Counter(str(r.get("Next Follow-up Date") or "").strip() for r in forms)

    res = subprocess.run([sys.executable, os.path.join(HERE, "build_dashboard.py")],
                         capture_output=True, text=True)
    if res.returncode != 0:
        print(res.stdout[-2000:], res.stderr[-2000:])
        stop("dashboard build failed")

    summary = {
        "week": f"{monday:%d %b} to {sunday:%d %b %Y}",
        "dates_in_file": f"{lo} to {hi}",
        "layout": layout,
        "checkin_rows": len(rows),
        "forms": len(forms),
        "unparsed_notes": len(unparsed),
        "reps_with_forms": len({r['Name'] for r in forms}),
        "gap_warning": gap,
        "new_dropdown_values": new_values,
        "top_follow_up_entries": fu.most_common(8),
        "cleaned_file": out,
        "dashboard": os.path.join(HERE, "index.html"),
    }
    print("SUMMARY " + json.dumps(summary, ensure_ascii=False, indent=1, default=str))


if __name__ == "__main__":
    main()
